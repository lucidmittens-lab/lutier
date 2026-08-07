"""
Builds LUTier_Keyboard_Shortcuts.xlsx -- a fillable form listing every
menu-bar dropdown item in LUTier (File, Edit, Workspace, Image Library,
LUT Library, CDL Library, Stack Memories, Viewer, Stack, Scopes, Graph,
Help), enumerated directly from main_window.py's _build_menu_bar and
each panel's populate_menu() (see stack_panel.py, sample_image_panel.py,
scopes_panel.py, lut_viewer_3d.py, image_library_panel.py,
lut_library_panel.py, cdl_library_panel.py, stack_memories_panel.py) --
not guessed or reconstructed from memory.

Tom fills in the "New Shortcut" column (yellow), sends the file back,
and that becomes the wiring list for QKeySequence/setShortcut calls in
the app.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Read Me
ws0 = wb.active
ws0.title = "Read Me"
ws0.sheet_view.showGridLines = False

title_font = Font(name=FONT_NAME, size=16, bold=True)
body_font = Font(name=FONT_NAME, size=11)
bold_body = Font(name=FONT_NAME, size=11, bold=True)

lines = [
    ("LUTier Keyboard Shortcuts -- Fill-In Form", title_font),
    ("", body_font),
    ("What this is:", bold_body),
    ("Every item in every dropdown menu in LUTier's menu bar (File, Edit, "
     "Workspace, Image Library, LUT Library, CDL Library, Stack Memories, "
     "Viewer, Stack, Scopes, Graph, Help), pulled directly from the app's "
     "code -- not reconstructed from memory.", body_font),
    ("", body_font),
    ("What to do:", bold_body),
    ("Go to the 'Menu Shortcuts' tab. Fill in the yellow 'New Shortcut' "
     "column for any item you want a keystroke for. Leave blank to skip "
     "an item. Send the file back and it'll be wired into the app.",
     body_font),
    ("", body_font),
    ("Format:", bold_body),
    ("Write shortcuts like: Ctrl+K, Ctrl+Shift+K, Alt+K, Ctrl+Alt+K. "
     "'Ctrl' automatically becomes 'Cmd' on macOS -- that's how the app's "
     "existing shortcuts are already written in the code, so stay "
     "consistent and use 'Ctrl', not 'Cmd', even though you'll press Cmd.",
     body_font),
    ("Example: to give \"Manual\" (Help menu) the keystroke Cmd+/, you'd "
     "write:  Ctrl+/", body_font),
    ("", body_font),
    ("Rows already showing a shortcut in gray are EXISTING shortcuts the "
     "app already has (File/Edit/Workspace menus) -- shown for reference "
     "so you don't accidentally pick a conflicting key. You can still "
     "type a replacement in the yellow column if you want to change one.",
     body_font),
    ("", body_font),
    ("A few items are grouped choices (e.g. Scopes: Histogram/Waveform/"
     "Vectorscope, or Graph: 3D/1D Y/1D R/1D G/1D B) -- each still gets "
     "its own row and can take its own independent shortcut.", body_font),
    ("", body_font),
    ("Items inside a submenu (e.g. Viewer > Left Image > Sample Before) "
     "show the submenu name in its own column so it's clear which group "
     "they belong to.", body_font),
]
row = 1
for text, font in lines:
    cell = ws0.cell(row=row, column=1, value=text)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    row += 1
ws0.column_dimensions["A"].width = 110
for r in range(1, row):
    ws0.row_dimensions[r].height = 18 if ws0.cell(row=r, column=1).value else 6

# ---------------------------------------------------------------- data
# (menu, submenu-or-blank, item label, type, current shortcut-or-blank, notes)
rows = [
    ("File", "", "New Session", "Action", "Ctrl+N", ""),
    ("File", "", "Open Session...", "Action", "Ctrl+O", ""),
    ("File", "", "Save Session", "Action", "Ctrl+S", ""),
    ("File", "", "Save Session As...", "Action", "Ctrl+Shift+S", ""),
    ("File", "", "Export...", "Action", "Ctrl+E", ""),

    ("Edit", "", "Undo", "Action", "Ctrl+Z", ""),
    ("Edit", "", "Redo", "Action", "Ctrl+Shift+Z", ""),

    ("Workspace", "", "Browser (show/hide)", "Toggle", "Shift+1", ""),
    ("Workspace", "", "Image Library (show/hide)", "Toggle", "Shift+2", ""),
    ("Workspace", "", "LUT Library (show/hide)", "Toggle", "Shift+3", ""),
    ("Workspace", "", "CDL Library (show/hide)", "Toggle", "Shift+4", ""),
    ("Workspace", "", "Stack Memories (show/hide)", "Toggle", "Shift+5", ""),
    ("Workspace", "", "Viewer (show/hide)", "Toggle", "Shift+6", ""),
    ("Workspace", "", "Stack (show/hide)", "Toggle", "Shift+7", ""),
    ("Workspace", "", "Parameters (show/hide)", "Toggle", "Shift+8", ""),
    ("Workspace", "", "Scopes (show/hide)", "Toggle", "Shift+9", ""),
    ("Workspace", "", "Graph (show/hide)", "Toggle", "Shift+0", ""),
    ("Workspace", "", "Float Browser", "Toggle", "Ctrl+Shift+1", ""),
    ("Workspace", "", "Float Image Library", "Toggle", "Ctrl+Shift+2", ""),
    ("Workspace", "", "Float LUT Library", "Toggle", "Ctrl+Shift+3", ""),
    ("Workspace", "", "Float CDL Library", "Toggle", "Ctrl+Shift+4", ""),
    ("Workspace", "", "Float Stack Memories", "Toggle", "Ctrl+Shift+5", ""),
    ("Workspace", "", "Float Viewer", "Toggle", "Ctrl+Shift+6", ""),
    ("Workspace", "", "Float Stack", "Toggle", "Ctrl+Shift+7", ""),
    ("Workspace", "", "Float Parameters", "Toggle", "Ctrl+Shift+8", ""),
    ("Workspace", "", "Float Scopes", "Toggle", "Ctrl+Shift+9", ""),
    ("Workspace", "", "Float Graph", "Toggle", "Ctrl+Shift+0", ""),
    ("Workspace", "", "Reset to Default Workspace", "Action", "", ""),
    ("Workspace", "", "Export Workspace Preset...", "Action", "", ""),
    ("Workspace", "", "Import Workspace Preset...", "Action", "", ""),

    ("Image Library", "", "Import...", "Action", "", ""),
    ("Image Library", "", "Load as Sample", "Action", "", ""),
    ("Image Library", "", "Load as Reference", "Action", "", ""),
    ("Image Library", "", "Rename...", "Action", "", ""),
    ("Image Library", "", "Delete", "Action", "", ""),

    ("LUT Library", "", "Import...", "Action", "", ""),
    ("LUT Library", "", "Apply to Stack", "Action", "", ""),
    ("LUT Library", "", "Rename...", "Action", "", ""),
    ("LUT Library", "", "Delete", "Action", "", ""),

    ("CDL Library", "", "Import...", "Action", "", ""),
    ("CDL Library", "", "Apply to Stack", "Action", "", ""),
    ("CDL Library", "", "Rename...", "Action", "", ""),
    ("CDL Library", "", "Delete", "Action", "", ""),

    ("Stack Memories", "", "Save Current Stack...", "Action", "", ""),
    ("Stack Memories", "", "Load", "Action", "", ""),
    ("Stack Memories", "", "Update", "Action", "", ""),
    ("Stack Memories", "", "Rename...", "Action", "", ""),
    ("Stack Memories", "", "Delete", "Action", "", ""),
    ("Stack Memories", "", "Export...", "Action", "", ""),
    ("Stack Memories", "", "Import...", "Action", "", ""),

    ("Viewer", "", "Load Sample Image...", "Action", "", ""),
    ("Viewer", "", "Adapt Sample Image...", "Action", "", ""),
    ("Viewer", "", "Export Still...", "Action", "", ""),
    ("Viewer", "", "Load Default Sample", "Action", "", ""),
    ("Viewer", "", "Load Reference Image...", "Action", "", ""),
    ("Viewer", "", "Adapt Reference...", "Action", "", ""),
    ("Viewer", "", "Enable Wipe", "Toggle", "", ""),
    ("Viewer", "Left Image", "Sample Before", "Choice", "", "exclusive w/ the other 2 rows in Left Image"),
    ("Viewer", "Left Image", "Sample After", "Choice", "", "exclusive w/ the other 2 rows in Left Image"),
    ("Viewer", "Left Image", "Reference", "Choice", "", "exclusive w/ the other 2 rows in Left Image"),
    ("Viewer", "Right Image", "Sample Before", "Choice", "", "exclusive w/ the other 2 rows in Right Image"),
    ("Viewer", "Right Image", "Sample After", "Choice", "", "exclusive w/ the other 2 rows in Right Image"),
    ("Viewer", "Right Image", "Reference", "Choice", "", "exclusive w/ the other 2 rows in Right Image"),

    ("Stack", "", "Add Printer Lights", "Action", "", ""),
    ("Stack", "", "Add ACES DRT", "Action", "", ""),
    ("Stack", "", "Add ASC CDL", "Action", "", ""),
    ("Stack", "", "Add Color Space Conversion", "Action", "", ""),
    ("Stack", "", "Add Fork (Split + Merge)", "Action", "", ""),
    ("Stack", "", "Edit Size/Domain...", "Action", "", ""),

    ("Scopes", "", "Histogram", "Choice", "", "exclusive w/ Waveform, Vectorscope"),
    ("Scopes", "", "Waveform", "Choice", "", "exclusive w/ Histogram, Vectorscope"),
    ("Scopes", "", "Vectorscope", "Choice", "", "exclusive w/ Histogram, Waveform"),
    ("Scopes", "Waveform Style", "RGB Parade", "Choice", "", "exclusive w/ the other 4 rows in Waveform Style"),
    ("Scopes", "Waveform Style", "Luma", "Choice", "", "exclusive w/ the other 4 rows in Waveform Style"),
    ("Scopes", "Waveform Style", "R", "Choice", "", "exclusive w/ the other 4 rows in Waveform Style"),
    ("Scopes", "Waveform Style", "G", "Choice", "", "exclusive w/ the other 4 rows in Waveform Style"),
    ("Scopes", "Waveform Style", "B", "Choice", "", "exclusive w/ the other 4 rows in Waveform Style"),
    ("Scopes", "Waveform Graticule", "None", "Choice", "", "exclusive w/ the other 3 rows in Waveform Graticule"),
    ("Scopes", "Waveform Graticule", "PQ", "Choice", "", "exclusive w/ the other 3 rows in Waveform Graticule"),
    ("Scopes", "Waveform Graticule", "mV", "Choice", "", "exclusive w/ the other 3 rows in Waveform Graticule"),
    ("Scopes", "Waveform Graticule", "%", "Choice", "", "exclusive w/ the other 3 rows in Waveform Graticule"),
    ("Scopes", "", "Graticule Intensity...", "Action", "", ""),

    ("Graph", "", "3D", "Choice", "", "exclusive w/ 1D Y, 1D R, 1D G, 1D B"),
    ("Graph", "", "1D Y", "Choice", "", "exclusive w/ 3D, 1D R, 1D G, 1D B"),
    ("Graph", "", "1D R", "Choice", "", "exclusive w/ 3D, 1D Y, 1D G, 1D B"),
    ("Graph", "", "1D G", "Choice", "", "exclusive w/ 3D, 1D Y, 1D R, 1D B"),
    ("Graph", "", "1D B", "Choice", "", "exclusive w/ 3D, 1D Y, 1D R, 1D G"),

    ("Help", "", "Manual", "Action", "", ""),
]

ws = wb.create_sheet("Menu Shortcuts")
ws.sheet_view.showGridLines = False

headers = ["Menu", "Submenu", "Item", "Type", "Current Shortcut", "New Shortcut", "Notes"]
header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="404040")
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = border
ws.freeze_panes = "A2"

data_font = Font(name=FONT_NAME, size=11)
menu_font = Font(name=FONT_NAME, size=11, bold=True)
existing_fill = PatternFill("solid", fgColor="E8E8E8")
existing_font = Font(name=FONT_NAME, size=11, color="666666")
fill_yellow = PatternFill("solid", fgColor="FFFF99")
notes_font = Font(name=FONT_NAME, size=10, italic=True, color="808080")

r = 2
prev_menu = None
for menu, submenu, item, type_, current, notes in rows:
    is_new_menu_group = (menu != prev_menu)
    prev_menu = menu

    ws.cell(row=r, column=1, value=menu).font = menu_font
    ws.cell(row=r, column=2, value=submenu).font = data_font
    ws.cell(row=r, column=3, value=item).font = data_font
    ws.cell(row=r, column=4, value=type_).font = data_font

    cur_cell = ws.cell(row=r, column=5, value=current if current else "—")
    if current:
        cur_cell.font = existing_font
        cur_cell.fill = existing_fill
    else:
        cur_cell.font = existing_font

    new_cell = ws.cell(row=r, column=6, value=None)
    new_cell.fill = fill_yellow
    new_cell.font = data_font

    notes_cell = ws.cell(row=r, column=7, value=notes)
    notes_cell.font = notes_font

    for c in range(1, 8):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=False)

    # Faint top border to visually separate menu groups
    if is_new_menu_group and r > 2:
        for c in range(1, 8):
            existing = ws.cell(row=r, column=c).border
            ws.cell(row=r, column=c).border = Border(
                left=existing.left, right=existing.right,
                top=Side(style="thin", color="808080"), bottom=existing.bottom)

    r += 1

widths = {"A": 16, "B": 20, "C": 32, "D": 10, "E": 16, "F": 16, "G": 42}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

for row_dim in range(1, r):
    ws.row_dimensions[row_dim].height = 16

wb.save("/sessions/compassionate-optimistic-brown/mnt/outputs/lut_tool/shortcuts_form/LUTier_Keyboard_Shortcuts.xlsx")
print(f"saved, {r - 2} data rows")
